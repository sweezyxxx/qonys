"""Assemble the agreed English analysis-stage report and editable DOCX."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / 'docs' / 'submission-2026-10-24'
TITLE = ('Development of an AI-Powered Residential Real Estate Visualization and Intelligent Property Selection '
         'Platform with 3D Site Visualization, Floor-Level Layouts, and Interactive Mapping')

def content(name):
    return (BASE / name).read_text(encoding='utf-8').strip()

def without_title(text):
    return '\n'.join(text.splitlines()[1:]).strip()

def part(number, title, text):
    body = without_title(text)
    return f'# {number} {title}\n\n{body}\n'

proposal = content('02-project-proposal.md')
proposal = re.sub(r'## Project title\n\n.*?\n\n', '', proposal, count=1, flags=re.S)
proposal = proposal.replace('Selecting an extra variant retains its context in comparison and enquiries.',
    'A compact dropdown beneath the price selects among offered conditions; a single condition is shown as text. '
    'Initially select the developer base unless a valid explicit comparison/AI context selects another offer. '
    'The selection retains its context in comparison and enquiries.')
proposal = proposal.replace('Email-delivery details will be specified separately.',
    'Buyer profile name, surname and phone are optional before an enquiry needs them, editable in settings '
    'and prefilled/correctable in the enquiry. Password can be changed; email is displayed without an address-change '
    'workflow. Signing in preserves the guest comparison list.')

external = content('26-external-evidence-records.md')
relevance = external.split('## Supervisor feedback')[0]
supervision = '# Supervisor feedback\n\n' + external.split('## Supervisor feedback', 1)[1].strip()
architecture = content('16-architecture.md')
architecture = re.sub(r'## Editable diagrams and readable exports\n.*?(?=## Logical module responsibilities)',
    '## Architecture figures\n\nEditable Mermaid/PlantUML sources, SVG/JSON geometry and PNG exports accompany the report '
    'in the diagrams directory. The semantic sources and exported geometry are maintained together. '
    'Figures show the system boundaries, actors and component responsibilities.\n\n', architecture, flags=re.S)
# The same technical risks and configuration dependencies are covered in section 10.
architecture = architecture.split('## Design risks and deferred configuration')[0].strip()

intro = f'''# {TITLE}

Analysis and Design Submission

24 October 2026

Team: Temirlan Almukhamedov, Yevgeniy Averyanov and Bexultan Nishanbek

Supervisor: Islam Omirzak

## Submission overview

QONYS is designed to combine a shared residential-development catalogue, physical-apartment selection, interactive floor layouts, infrastructure maps and optional developer-supplied visual assets. Its AI functions are limited to catalogue-grounded buyer preference interpretation and correctable image-region suggestions. Developers publish their own complete reviewed developments, and administrators manage developer access and public restrictions.

The analysis and design milestone produced an agreed scope, findings from three interviews, literature and system comparisons, 40 functional and 10 non-functional requirements with acceptance criteria, architectural diagrams and an interactive scenario prototype. A verification protocol and implementation plan connect these design decisions to the next project stage. The report presents the completed analysis and design work; its application behaviour is specified by the requirements and prototype.

## Contents

1. Project proposal
2. Evidence of relevance
3. Literature review
4. Existing system comparison
5. Technology comparison
6. Requirements specification
7. Architecture and diagrams
8. UX and prototype
9. Verification and evaluation plan
10. Project management and risks
11. Supervisor feedback

'''

sections = [
    part(1, 'Project proposal', proposal),
    part(2, 'Evidence of relevance', relevance),
    part(3, 'Literature review', content('22-literature-review.md')),
    part(4, 'Existing system comparison', content('23-existing-system-comparison.md')),
    part(5, 'Technology comparison', content('24-technology-comparison.md')),
    part(6, 'Requirements specification', content('15-requirements.md')),
    part(7, 'Architecture and diagrams', architecture),
    part(8, 'UX and prototype', content('27-prototype-guide.md')),
    part(9, 'Verification and evaluation plan', content('17-evaluation-plan.md')),
    part(10, 'Project management and risks', content('25-management-and-risks.md')),
    part(11, 'Supervisor feedback', supervision),
]

figure_blocks = {
    7: [('System context', 'diagrams/01-system-context.png'),
        ('Use cases', 'diagrams/02-use-cases.png'),
        ('Component architecture', 'diagrams/03-components.png')],
    8: [('Apartment and compact finishing selection', 'prototype/report-screens/04-apartment.png'),
        ('Draft saving and publication validation', 'prototype/report-screens/07-publication.png'),
        ('Administrator hiding reason and restoration', 'prototype/report-screens/09-administrator.png')]
}

full = intro
for number, body in enumerate(sections, 1):
    full += body + '\n\n'
    for label, path in figure_blocks.get(number, []):
        full += f'![{label}]({path})\n\n'
(BASE / 'QONYS_analysis_design_2026-10-24.md').write_text(full.rstrip()+'\n', encoding='utf-8')

doc = Document()
normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08
for name, size in [('Title', 20), ('Subtitle', 13), ('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 11.5), ('Heading 4', 11)]:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(14 if name.startswith('Heading') else 0)
    style.paragraph_format.space_after = Pt(8)
    style.font.underline = False
    style.font.italic = False
    for border in style._element.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)
    for fonts in style._element.xpath('./w:rPr/w:rFonts'):
        for attr in list(fonts.attrib):
            if attr.endswith('Theme') or attr.endswith('theme'):
                del fonts.attrib[attr]

def setup_section(section, landscape=False):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Inches(11.69 if landscape else 8.27)
    section.page_height = Inches(8.27 if landscape else 11.69)
    section.top_margin = Inches(.65)
    section.bottom_margin = Inches(.65)
    section.left_margin = Inches(.7)
    section.right_margin = Inches(.7)
    section.header_distance = Inches(.25)
    section.footer_distance = Inches(.3)

setup_section(doc.sections[0])
footer = doc.sections[0].footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run('QONYS  |  ')
run.font.size = Pt(9)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def section(landscape=False):
    s = doc.add_section(WD_SECTION_START.NEW_PAGE)
    setup_section(s, landscape)
    return s

inline = re.compile(r'\[([^\]]+)\]\(([^)]+)\)|\*\*([^*]+)\*\*|`([^`]+)`|\*([^*]+)\*')

def rich(paragraph, text, size=None):
    pos = 0
    for match in inline.finditer(text):
        paragraph.add_run(text[pos:match.start()])
        if match[1] is not None:
            target = match[2]
            hyperlink = OxmlElement('w:hyperlink')
            rel = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
            hyperlink.set(qn('r:id'), rel)
            r = OxmlElement('w:r')
            props = OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'), '145B99'); props.append(color)
            u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); props.append(u)
            r.append(props)
            if size:
                sz=OxmlElement('w:sz'); sz.set(qn('w:val'),str(round(size*2))); props.append(sz)
            t = OxmlElement('w:t'); t.text = match[1]; r.append(t)
            hyperlink.append(r); paragraph._p.append(hyperlink)
        else:
            run = paragraph.add_run(next(x for x in match.groups()[2:] if x is not None))
            run.bold = match[3] is not None
            run.italic = match[5] is not None
            if match[4] is not None:
                run.font.name = 'Consolas'
        pos = match.end()
    paragraph.add_run(text[pos:])
    if size:
        for r in paragraph.runs:
            r.font.size = Pt(size)

def add_table(lines):
    rows = [[c.strip() for c in line.strip().strip('|').split('|')] for line in lines]
    rows = [r for r in rows if not all(re.fullmatch(r'[:\- ]+', c or '-') for c in r)]
    columns = len(rows[0])
    table = doc.add_table(rows=0, cols=columns)
    table.autofit = False
    available = doc.sections[-1].page_width.inches - 1.4
    first = rows[0][0]
    if first == 'ID' and columns == 4 and rows[0][1] == 'Scientific source':
        weights = [.045, .30, .24, .415]
    elif first == 'ID' and columns == 4:
        weights = [.07, .30, .10, .53]
    elif columns == 5:
        weights = [.18, .205, .205, .205, .205]
    elif columns == 4:
        if first == 'Risk':
            weights = [.17, .21, .43, .19]
        elif first == 'Local planning ID':
            weights = [.12, .32, .23, .33]
        elif first == 'Milestone':
            weights = [.19, .21, .38, .22]
        else:
            weights = [.20, .27, .28, .25]
    elif columns == 3:
        if first in ('Participant code', 'Code'):
            weights = [.09, .59, .32]
        elif first == 'Finding and reported experience':
            weights = [.36, .22, .42]
        elif first == 'Candidate need and basis':
            weights = [.30, .47, .23]
        elif first == 'ID':
            weights = [.055, .39, .555]
        elif first == 'Design question':
            weights = [.30, .14, .56]
        elif first == 'Check':
            weights = [.24, .57, .19]
        else:
            weights = [.23, .38, .39]
    else:
        weights = [.3, .7] if columns == 2 else [1/columns]*columns
    for col, weight in zip(table.columns, weights):
        col.width = Inches(available*weight)
    borders = OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        b = OxmlElement('w:'+side); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'5'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
    table._tbl.tblPr.append(borders)
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for j, cell in enumerate(cells):
            cell.width = Inches(available*weights[j])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ('top','bottom','left','right'):
                x=OxmlElement('w:'+side); x.set(qn('w:w'),'95'); x.set(qn('w:type'),'dxa'); margins.append(x)
            props.append(margins)
            if index == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'),'E6EFF7'); props.append(shade)
            p=cell.paragraphs[0]
            p.paragraph_format.space_after=Pt(3)
            p.paragraph_format.space_before=Pt(3)
            p.paragraph_format.line_spacing=1.05
            p.paragraph_format.keep_with_next=index == 0
            if (first=='ID' and j==0) or rows[0][j]=='Priority':
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            rich(p, row[j] if j < len(row) else '', 10)
            for r in p.runs:
                r.bold = index == 0 or r.bold
        trpr = table.rows[-1]._tr.get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit'); trpr.append(no_split)
        if index == 0:
            repeat=OxmlElement('w:tblHeader'); trpr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def new_numbering():
    numbering=doc.part.numbering_part.element
    abstract_ids=[int(x.get(qn('w:abstractNumId'))) for x in numbering.findall(qn('w:abstractNum'))]
    abstract_id=max(abstract_ids,default=-1)+1
    abstract=OxmlElement('w:abstractNum'); abstract.set(qn('w:abstractNumId'),str(abstract_id))
    level=OxmlElement('w:lvl'); level.set(qn('w:ilvl'),'0')
    for tag,value in [('start','1'),('numFmt','decimal'),('lvlText','%1.'),('lvlJc','left')]:
        x=OxmlElement('w:'+tag); x.set(qn('w:val'),value); level.append(x)
    pp=OxmlElement('w:pPr'); ind=OxmlElement('w:ind'); ind.set(qn('w:left'),'360'); ind.set(qn('w:hanging'),'360'); pp.append(ind); level.append(pp)
    abstract.append(level); numbering.append(abstract)
    num=numbering.add_num(abstract_id)
    return num.numId

def markdown(text, skip_first=False, bookmark=None):
    lines=text.splitlines()
    if skip_first: lines=lines[1:]
    i=0
    while i < len(lines):
        line=lines[i].strip()
        if not line: i+=1; continue
        if line.startswith('|'):
            table=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                table.append(lines[i]); i+=1
            add_table(table); continue
        heading=re.match(r'^(#{1,6}) (.+)',line)
        if heading:
            level=min(len(heading[1]),4)
            text=heading[2].replace(':','').replace(' / ',' and ')
            p=doc.add_heading(text,level=level)
            if text in ('Uploaded asset storage', 'Image region suggestions', 'Project coordination and version control'):
                p.paragraph_format.page_break_before=True
            if level==1 and bookmark is not None:
                start=OxmlElement('w:bookmarkStart'); start.set(qn('w:id'),str(bookmark)); start.set(qn('w:name'),f'section_{bookmark}')
                end=OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'),str(bookmark))
                p._p.insert(1 if p._p.pPr is not None else 0,start); p._p.append(end)
        elif line.startswith('- '):
            rich(doc.add_paragraph(style='List Bullet'),line[2:])
        elif re.match(r'^\d+\. ',line):
            num_id=new_numbering()
            while i<len(lines):
                current=lines[i].strip()
                if not current:
                    i+=1; continue
                if not re.match(r'^\d+\. ',current):
                    break
                p=doc.add_paragraph()
                np=p._p.get_or_add_pPr().get_or_add_numPr()
                np.get_or_add_ilvl().val=0; np.get_or_add_numId().val=num_id
                rich(p,re.sub(r'^\d+\. ','',current))
                i+=1
            continue
        else:
            p = doc.add_paragraph()
            rich(p,line)
            next_line = next((x.strip() for x in lines[i+1:] if x.strip()), '')
            if next_line.startswith('|'):
                p.paragraph_format.keep_with_next = True
        i+=1

p=doc.add_paragraph(style='Title')
p.paragraph_format.space_before=Pt(26)
p.add_run(TITLE)
p=doc.add_paragraph('Analysis and Design Submission',style='Subtitle')
p.paragraph_format.space_after=Pt(28)
doc.add_paragraph('24 October 2026')
doc.add_paragraph('Temirlan Almukhamedov\nYevgeniy Averyanov\nBexultan Nishanbek')
doc.add_paragraph('Supervisor  Islam Omirzak')
doc.add_heading('Submission overview',level=2)
overview=intro.split('## Submission overview\n\n',1)[1].split('## Contents')[0].strip()
markdown(overview)
section()
doc.add_heading('Contents',1)
for n,title in enumerate(['Project proposal','Evidence of relevance','Literature review','Existing system comparison','Technology comparison','Requirements specification','Architecture and diagrams','UX and prototype','Verification and evaluation plan','Project management and risks','Supervisor feedback'],1):
    p=doc.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.8),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
    link=OxmlElement('w:hyperlink'); link.set(qn('w:anchor'),f'section_{n}')
    run=OxmlElement('w:r'); text=OxmlElement('w:t'); text.text=f'{n}  {title}'; run.append(text); link.append(run); p._p.append(link)
    toc_file=BASE/'contents-pages.json'
    if toc_file.exists():
        import json
        page_map=json.loads(toc_file.read_text(encoding='utf-8'))
        p.add_run('\t'+str(page_map.get(str(n),'')))
for n,text in enumerate(sections,1):
    if n != 11:
        section(landscape=n in (3,4,6))
    markdown(text,bookmark=n)
    if n in figure_blocks:
        for label, filename in figure_blocks[n]:
            section(landscape=True)
            doc.add_heading(label,level=2)
            image_path=BASE/filename
            from PIL import Image
            with Image.open(image_path) as im:
                ratio=im.height/im.width
            width=min(10.15,(6.0 if n==8 else 5.65)/ratio)
            p=doc.add_paragraph()
            p.paragraph_format.space_after=Pt(5)
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(); r.add_picture(str(image_path),width=Inches(width))
            inline_shape=r._r.xpath('.//wp:docPr')[0]
            inline_shape.set('descr',label+' for the planned QONYS system')
            caption=doc.add_paragraph()
            rich(caption, 'QONYS architecture design.' if n==7 else 'QONYS interface design using fictional example data.', 10)

doc.core_properties.title=TITLE
doc.core_properties.subject='Analysis and design submission for 24 October 2026'
doc.core_properties.author='QONYS project team'
doc.core_properties.keywords='QONYS, requirements, architecture, prototype, analysis'
output=BASE/'QONYS_analysis_design_2026-10-24.docx'
doc.save(output)
print(output)
print(f'{len(doc.paragraphs)} paragraphs; {len(doc.tables)} tables; {len(doc.inline_shapes)} figures')
